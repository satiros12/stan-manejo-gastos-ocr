#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extrae datos de extractos bancarios: .xlsx, .xlsm, .csv, .txt, .pdf (con texto).

Uso:  python3 import_extract.py <archivo>
Salida: JSON en stdout:
  {"kind":"table","headers":[...],"rows":[[...]],"total":N}   (Excel/CSV)
  {"kind":"lines","lines":[...],"total":N}                    (PDF con texto)
  {"kind":"scanned", ...} | {"error": "..."}
"""
import sys, os, json, csv, io, subprocess, datetime
from typing import NoReturn

MAX_ROWS = 3000
MAX_LINES = 1500
TERMS = ("fecha", "date", "concepto", "descrip", "importe", "amount", "cargo",
         "abono", "saldo", "movimiento", "detalle", "operacion", "operación",
         "comercio", "beneficiario", "debe", "haber")


def emit(obj) -> NoReturn:
    print(json.dumps(obj, ensure_ascii=False, default=str))
    sys.exit(0)


def norm_cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return v


def find_header(rows):
    for i, r in enumerate(rows[:25]):
        cells = [str(c).lower() for c in r if str(c).strip()]
        if len(cells) >= 2:
            joined = " ".join(cells)
            if any(t in joined for t in TERMS):
                return i
    for i, r in enumerate(rows[:25]):
        if len([c for c in r if str(c).strip()]) >= 2:
            return i
    return 0


def table_result(rows, sheet=None):
    rows = rows[:MAX_ROWS]
    if not rows:
        emit({"error": "El archivo no contiene filas"})
    h = find_header(rows)
    header_row = rows[h]
    data = rows[h + 1:]
    width = max([len(header_row)] + [len(r) for r in data]) if data else len(header_row)
    headers = [str(c) if str(c).strip() else f"col{j}" for j, c in enumerate(header_row)]
    headers += [f"col{j}" for j in range(len(headers), width)]
    data = [list(r) + [""] * (width - len(r)) for r in data]
    data = [r for r in data if any(str(c).strip() for c in r)]
    return {"kind": "table", "sheet": sheet, "header_row": h,
            "headers": headers[:width], "rows": data, "total": len(data)}


def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    best = None  # (n_filas_no_vacias, nombre, filas)
    for ws in wb.worksheets[:5]:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([norm_cell(v) for v in row])
            if len(rows) >= MAX_ROWS + 30:
                break
        n = sum(1 for r in rows if any(str(c).strip() for c in r))
        if best is None or n > best[0]:
            best = (n, ws.title, rows)
        if n > 20:
            break  # primera hoja con datos claros
    wb.close()
    if not best:
        emit({"error": "El Excel no tiene hojas con datos"})
    return table_result(best[2], best[1])


def read_csv(path):
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        emit({"error": "No se pudo decodificar el CSV"})
    sample = text[:8192]
    delim = max([";", ",", "\t", "|"], key=lambda d: sample.count(d))
    rd = csv.reader(io.StringIO(text), delimiter=delim)
    rows = []
    for i, r in enumerate(rd):
        if i >= MAX_ROWS + 30:
            break
        rows.append([norm_cell(c) for c in r])
    return table_result(rows)


def read_pdf(path):
    try:
        r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", path, "-"],
                           capture_output=True, text=True, timeout=90)
    except FileNotFoundError:
        emit({"error": "pdftotext no está instalado"})
    if r.returncode != 0:
        emit({"error": f"pdftotext falló: {(r.stderr or '')[:200]}"})
    text = r.stdout or ""
    lines = [ln.rstrip() for ln in text.splitlines()]
    lines = [ln for ln in lines if len(ln.strip()) >= 3]
    if len("".join(lines)) < 80:
        emit({"kind": "scanned",
              "hint": "El PDF no tiene capa de texto (parece escaneado). Usa el flujo de OCR/foto."})
    return {"kind": "lines", "lines": lines[:MAX_LINES], "total": len(lines)}


def main():
    if len(sys.argv) < 2:
        emit({"error": "uso: import_extract.py <archivo>"})
    p = sys.argv[1]
    if not os.path.exists(p):
        emit({"error": f"no existe: {p}"})
    ext = os.path.splitext(p)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        emit(read_xlsx(p))
    if ext == ".xls":
        emit({"error": "El formato .xls antiguo no está soportado; guarda el archivo como .xlsx o CSV"})
    if ext in (".csv", ".txt"):
        emit(read_csv(p))
    if ext == ".pdf":
        emit(read_pdf(p))
    emit({"error": f"extensión no soportada: {ext}"})


if __name__ == "__main__":
    main()
